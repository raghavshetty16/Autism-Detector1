from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = BASE_DIR / "Autism_Detection_Project_Report.docx"
WORK_DIR = BASE_DIR / "report_assets"
WORK_DIR.mkdir(exist_ok=True)

NAVY = "183B4E"
TEAL = "0E766E"
TEAL_LIGHT = "E8F3F1"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202A2E"
MUTED = "5F6B70"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D7DEE2"
WHITE = "FFFFFF"
GOLD = "B3812D"
RED = "9B1C1C"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margins=CELL_MARGINS):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in margins.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError("Table widths must total 9360 DXA")
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_table_text(table, header=True, font_size=9.2):
    for row_index, row in enumerate(table.rows):
        if header and row_index == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        color=WHITE if header and row_index == 0 else INK,
                        bold=header and row_index == 0,
                    )
        if header and row_index == 0:
            for cell in row.cells:
                set_cell_shading(cell, NAVY)
    set_table_borders(table)


def add_numbering_definition(doc, num_id, abstract_id, num_format, level_text):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), level_text)
    level.append(text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def add_list_item(doc, text, ordered=False, bold_prefix=None):
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "42" if ordered else "41")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(color)
    run_pr.append(underline)
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    set_run_font(run, size=9, color=MUTED)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and choose Update Field if the table of contents is not visible."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def add_toc_entry(doc, text, page="", level=0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22 * level)
    paragraph.paragraph_format.space_after = Pt(4 if level else 6)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.25),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    run = paragraph.add_run(text)
    set_run_font(run, size=10 if level else 10.5, color=INK, bold=level == 0)
    if page:
        page_run = paragraph.add_run(f"\t{page}")
        set_run_font(page_run, size=10, color=MUTED)


def add_caption(doc, text):
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=MUTED, italic=True)
    return paragraph


def add_body(doc, text, bold_lead=None, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = False
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, italic=italic)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, italic=italic)
    return paragraph


def add_callout(doc, title, text, color=TEAL, fill=TEAL_LIGHT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=color, size=10)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    title_run = paragraph.add_run(title.upper())
    set_run_font(title_run, size=9, color=color, bold=True)
    detail = cell.add_paragraph()
    detail.paragraph_format.space_after = Pt(0)
    detail.paragraph_format.line_spacing = 1.08
    detail_run = detail.add_run(text)
    set_run_font(detail_run, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=MID_GRAY, size=5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F8FA")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(code)
    set_run_font(run, name="Consolas", size=8.4, color=INK)


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def create_architecture_diagram(path):
    canvas = Image.new("RGB", (1800, 820), f"#{WHITE}")
    draw = ImageDraw.Draw(canvas)
    title_font = font(46, True)
    box_title = font(29, True)
    body_font = font(23)
    arrow_font = font(44, True)
    draw.text((80, 50), "Production prediction workflow", font=title_font, fill="#183B4E")
    boxes = [
        (70, 190, 370, 540, "Browser", ["Profile fields", "10 Yes/No answers", "Result modal"]),
        (430, 190, 730, 540, "FastAPI", ["Input validation", "AQ-10 scoring", "Jinja2 rendering"]),
        (790, 190, 1090, 540, "Model bundle", ["MinMaxScaler", "One-hot columns", "Linear SVC"]),
        (1150, 190, 1450, 540, "Explanation", ["Top contributions", "AQ point details", "Plain language"]),
        (1500, 190, 1730, 540, "MongoDB", ["Profile", "Answers", "Model output"]),
    ]
    colors = ["#E8F3F1", "#EAF1F7", "#FFF4DC", "#F2F4F7", "#E8F3F1"]
    outlines = ["#0E766E", "#2E74B5", "#B3812D", "#5F6B70", "#0E766E"]
    for index, (x1, y1, x2, y2, heading, lines) in enumerate(boxes):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=colors[index], outline=outlines[index], width=5)
        draw.text((x1 + 26, y1 + 35), heading, font=box_title, fill="#202A2E")
        y = y1 + 120
        for line in lines:
            draw.ellipse((x1 + 28, y + 8, x1 + 42, y + 22), fill=outlines[index])
            draw.text((x1 + 57, y), line, font=body_font, fill="#202A2E")
            y += 62
        if index < len(boxes) - 1:
            draw.text((x2 + 16, 325), ">", font=arrow_font, fill="#5F6B70")
    draw.rounded_rectangle((430, 640, 1450, 755), radius=18, fill="#183B4E")
    note = "The model is loaded once at startup. Requests perform preprocessing and prediction only."
    draw.text((500, 675), note, font=body_font, fill=f"#{WHITE}")
    canvas.save(path, quality=95)


def create_benchmark_chart(path):
    data = [
        ("Decision Tree", None, 1.000),
        ("Random Forest", 0.990, 1.000),
        ("Linear SVM", 1.000, 1.000),
        ("KNN (k=10)", 0.946, 0.918),
        ("Naive Bayes", 0.875, 0.768),
        ("Logistic Reg.", 0.997, 0.995),
    ]
    canvas = Image.new("RGB", (1700, 950), f"#{WHITE}")
    draw = ImageDraw.Draw(canvas)
    title_font = font(42, True)
    axis_font = font(23)
    label_font = font(24, True)
    small_font = font(19)
    draw.text((70, 40), "Notebook benchmark comparison", font=title_font, fill="#183B4E")
    draw.text((70, 102), "10-fold cross-validation mean vs. holdout F0.5 score", font=axis_font, fill="#5F6B70")
    x0, x1 = 380, 1600
    y0 = 190
    row_h = 104
    scale_min = 0.70
    for tick in (0.70, 0.80, 0.90, 1.00):
        x = x0 + int((tick - scale_min) / (1.0 - scale_min) * (x1 - x0))
        draw.line((x, y0 - 15, x, y0 + row_h * len(data)), fill="#D7DEE2", width=2)
        draw.text((x - 20, y0 + row_h * len(data) + 15), f"{tick:.2f}", font=small_font, fill="#5F6B70")
    for index, (name, cv, fscore) in enumerate(data):
        y = y0 + index * row_h
        draw.text((65, y + 23), name, font=label_font, fill="#202A2E")
        if cv is not None:
            cv_end = x0 + int((cv - scale_min) / (1.0 - scale_min) * (x1 - x0))
            draw.rounded_rectangle((x0, y + 11, cv_end, y + 43), radius=8, fill="#2E74B5")
            draw.text((min(cv_end + 10, x1 - 65), y + 8), f"{cv:.3f}", font=small_font, fill="#2E74B5")
        else:
            draw.text((x0 + 10, y + 8), "not reported", font=small_font, fill="#5F6B70")
        f_end = x0 + int((fscore - scale_min) / (1.0 - scale_min) * (x1 - x0))
        draw.rounded_rectangle((x0, y + 53, f_end, y + 85), radius=8, fill="#0E766E")
        draw.text((min(f_end + 10, x1 - 65), y + 50), f"{fscore:.3f}", font=small_font, fill="#0E766E")
    draw.rectangle((68, 865, 94, 891), fill="#2E74B5")
    draw.text((108, 862), "10-fold CV mean", font=small_font, fill="#202A2E")
    draw.rectangle((355, 865, 381, 891), fill="#0E766E")
    draw.text((395, 862), "Holdout F0.5", font=small_font, fill="#202A2E")
    canvas.save(path, quality=95)


def create_dataset_chart(path):
    canvas = Image.new("RGB", (1500, 620), f"#{WHITE}")
    draw = ImageDraw.Draw(canvas)
    title_font = font(40, True)
    number_font = font(54, True)
    label_font = font(23)
    draw.text((60, 40), "Dataset after cleaning", font=title_font, fill="#183B4E")
    cards = [
        (60, 150, 450, 500, "702", "usable records", "#EAF1F7", "#2E74B5"),
        (555, 150, 945, 500, "189", "YES class (26.9%)", "#FFF4DC", "#B3812D"),
        (1050, 150, 1440, 500, "513", "NO class (73.1%)", "#E8F3F1", "#0E766E"),
    ]
    for x1, y1, x2, y2, number, label, fill, outline in cards:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=fill, outline=outline, width=5)
        bbox = draw.textbbox((0, 0), number, font=number_font)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y1 + 75), number, font=number_font, fill=outline)
        lines = wrap(label, width=20)
        y = y1 + 185
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=label_font)
            draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=label_font, fill="#202A2E")
            y += 36
    canvas.save(path, quality=95)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    style_tokens = {
        "Title": (28, NAVY, 0, 8),
        "Subtitle": (14, TEAL, 0, 16),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in style_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name.startswith("Heading")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.color.rgb = rgb(MUTED)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(6)

    if "Source Text" not in styles:
        source_style = styles.add_style("Source Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source_style = styles["Source Text"]
    source_style.font.name = "Calibri"
    source_style.font.size = Pt(9)
    source_style.font.color.rgb = rgb(MUTED)
    source_style.paragraph_format.space_before = Pt(4)
    source_style.paragraph_format.space_after = Pt(4)

    add_numbering_definition(doc, 41, 41, "bullet", "•")
    add_numbering_definition(doc, 42, 42, "decimal", "%1.")

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hr = hp.add_run("AUTISM DETECTION USING MACHINE LEARNING  |  PROJECT REPORT")
    set_run_font(hr, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(2)
    fr = fp.add_run("Page ")
    set_run_font(fr, size=9, color=MUTED)
    add_page_field(fp)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "false")


def add_cover(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(12)
    kicker = paragraph.add_run("SYSTEM IMPLEMENTATION REPORT")
    set_run_font(kicker, size=10, color=GOLD, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Autism Trait Screening\nUsing Machine Learning")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("A FastAPI system with explainable results and MongoDB data collection")

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = image_paragraph.add_run().add_picture(str(BASE_DIR / "image.jpeg"), width=Inches(5.75))
    set_alt_text(shape, "Project image", "Child interacting with colorful blocks, used as the visual identity of the screening application.")

    metadata = doc.add_table(rows=4, cols=2)
    set_table_geometry(metadata, [2300, 7060])
    metadata_rows = [
        ("Course", "Introduction to AI and Generative AI"),
        ("Prepared by", "____________________________"),
        ("Student ID", "____________________________"),
        ("Submission date", "June 2026"),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_rows):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], TEAL_LIGHT)
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=10, color=TEAL, bold=True)
        for run in row.cells[1].paragraphs[0].runs:
            set_run_font(run, size=10, color=INK)
    set_table_borders(metadata, color=MID_GRAY, size=5)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(14)
    note_run = note.add_run("Screening support only - not a medical diagnosis")
    set_run_font(note_run, size=10.5, color=RED, bold=True)
    doc.add_page_break()


def add_front_matter(doc):
    doc.add_heading("Abstract", level=1)
    add_body(
        doc,
        "This project implements a web-based machine-learning system for screening autism-related traits using ten AQ-10 questions and selected profile information. The work begins with a 704-record adult autism screening dataset, removes incomplete records, scales numeric variables, one-hot encodes categorical variables, and compares six classification techniques. A Support Vector Classifier is optimized through grid search using an F-beta score with beta equal to 0.5. The selected linear SVC is serialized with its preprocessing information in a joblib bundle, allowing the FastAPI application to load the trained model once rather than retraining during startup or each request. Users complete a responsive questionnaire, receive a YES or NO screening result, see a plain-language explanation of influential inputs and AQ-10 scoring, and can optionally save responses to MongoDB Atlas for reviewed future training. The notebook reports perfect holdout accuracy and F0.5 for the optimized model. This result demonstrates strong fit to the available dataset but is interpreted cautiously because the dataset is small, class-imbalanced, and includes a derived AQ total score that is closely related to the target. The completed system demonstrates the full AI lifecycle from data preparation and benchmarking to model persistence, API integration, user interface design, explainability, and responsible deployment considerations."
    )
    keywords = doc.add_paragraph()
    key_run = keywords.add_run("Keywords: ")
    set_run_font(key_run, bold=True, color=TEAL)
    value = keywords.add_run("autism screening, AQ-10, machine learning, SVC, FastAPI, explainability, MongoDB")
    set_run_font(value)

    doc.add_heading("Academic Integrity and AI Use Declaration", level=1)
    add_body(
        doc,
        "This report documents the implementation contained in the submitted project files. AI assistance was used to support software development, wording, organization, and document formatting. The student remains responsible for verifying the implementation, understanding every technical choice, checking all results, citing sources, and explaining the work during evaluation. No model result in this report should be interpreted as a clinical diagnosis."
    )
    add_callout(
        doc,
        "Submission note",
        "The course guidelines require the final report in PDF format. This editable DOCX can be reviewed, completed with student details, and exported to PDF before submission.",
        color=GOLD,
        fill="FFF7E8",
    )

    doc.add_heading("Table of Contents", level=1)
    toc_entries = [
        ("Abstract", 0),
        ("Academic Integrity and AI Use Declaration", 0),
        ("1. Problem Definition", 0),
        ("2. Motivation and Objectives", 0),
        ("3. Background and Related Concepts", 0),
        ("4. System Architecture", 0),
        ("5. Dataset and Data Preparation", 0),
        ("6. Machine-Learning Methodology", 0),
        ("7. Implementation Details", 0),
        ("8. Benchmarking and Evaluation", 0),
        ("9. Results and Discussion", 0),
        ("10. Limitations and Responsible AI", 0),
        ("11. Future Improvements", 0),
        ("12. Conclusion", 0),
        ("References", 0),
        ("Appendix A. Project File Guide", 0),
        ("Appendix B. Run and Demonstration Guide", 0),
        ("Appendix C. Example API Request", 0),
        ("Appendix D. Q&A Preparation", 0),
    ]
    for entry, level in toc_entries:
        add_toc_entry(doc, entry, level=level)
    doc.add_page_break()


def add_problem_and_background(doc):
    doc.add_heading("1. Problem Definition", level=1)
    add_body(
        doc,
        "Autism assessment normally requires qualified clinical judgment, detailed history, observation, and validated instruments. Access to a complete evaluation may be delayed by time, cost, location, or limited awareness. A short digital screener can help a person organize answers and decide whether discussing concerns with a professional may be useful. The engineering problem is therefore to convert a research dataset and questionnaire into a consistent, understandable, and reusable screening workflow without presenting the output as a diagnosis."
    )
    add_body(
        doc,
        "The implemented system accepts seven profile fields and ten Yes/No AQ-10 responses. It converts the responses into the score representation expected by the saved model, applies the same transformations used during training, predicts a binary class, explains important inputs, and stores the completed response when MongoDB is available."
    )
    add_callout(
        doc,
        "Problem statement",
        "Build a usable AI screening application that reproduces the notebook model's preprocessing and prediction logic, avoids repeated training, explains the result in simple language, and safely collects future data without claiming clinical certainty.",
    )

    doc.add_heading("2. Motivation and Objectives", level=1)
    add_body(
        doc,
        "The project was selected because it combines a socially meaningful classification problem with practical AI engineering. It is not enough to obtain a score inside a notebook; a complete system must validate inputs, preserve model preprocessing, serve predictions through an interface, communicate uncertainty, and handle storage failures gracefully."
    )
    objectives = [
        "Study the adult autism screening dataset and understand the target and feature distributions.",
        "Compare multiple classical machine-learning classifiers rather than choosing one without evidence.",
        "Tune the final classifier using cross-validated grid search and a stated evaluation metric.",
        "Persist the trained model and transformations so real-time requests do not retrain the model.",
        "Create a responsive frontend that presents the ten real questions with simple Yes/No controls.",
        "Explain the prediction and the AQ-10 point calculation in plain language.",
        "Store completed responses in MongoDB for later review and possible retraining.",
        "Document limitations, privacy concerns, bias risks, and the difference between screening and diagnosis.",
    ]
    for objective in objectives:
        add_list_item(doc, objective)

    doc.add_heading("3. Background and Related Concepts", level=1)
    doc.add_heading("3.1 AQ-10 screening", level=2)
    add_body(
        doc,
        "The AQ-10 is a short form derived from the longer Autism-Spectrum Quotient. It uses ten items chosen for their ability to distinguish autistic traits in the original development work [1]. Answers are scored dichotomously, and some items are reverse-scored. Published work describes the instrument as a screener, not a replacement for diagnostic evaluation [1][2]. The present application simplifies the original response scale into Yes/No controls to match the project's dataset representation."
    )
    doc.add_heading("3.2 Supervised classification", level=2)
    add_body(
        doc,
        "Supervised classification learns a mapping from input features to known labels. In this project, the label `Class/ASD` is converted from YES/NO text to 1/0. The algorithms learn patterns from training rows and are then evaluated on held-out rows that were not used for fitting. The notebook compares tree-based, distance-based, probabilistic, linear, and margin-based approaches so that the selected technique has an empirical basis."
    )
    doc.add_heading("3.3 Why the system is screening support", level=2)
    add_body(
        doc,
        "A classifier can reproduce patterns in its training data, but it cannot establish a clinical diagnosis. Dataset labels may reflect prior screening rather than complete clinical assessment, and performance on a random split does not guarantee performance for new populations. The interface therefore uses careful wording and includes an explicit medical disclaimer."
    )


def add_architecture(doc, architecture_path):
    doc.add_heading("4. System Architecture", level=1)
    add_body(
        doc,
        "The application follows a small layered architecture. The browser is responsible for data entry and presentation. FastAPI handles routes, validation, AQ scoring, prediction orchestration, and template rendering. The serialized bundle owns the fitted model and preprocessing contract. The explanation layer translates numeric contributions into readable factors. The database layer saves responses without preventing a prediction when storage is unavailable."
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(architecture_path), width=Inches(6.35))
    set_alt_text(shape, "System architecture", "Flow from browser to FastAPI, saved model bundle, result explanation, and MongoDB storage.")
    add_caption(doc, "Figure 1. Production prediction workflow and component responsibilities.")

    doc.add_heading("4.1 Component responsibilities", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Technology"
    table.rows[0].cells[2].text = "Responsibility"
    rows = [
        ("Frontend", "HTML, CSS, JavaScript, Jinja2", "Collect answers, retain form values, show the result modal, and support responsive interaction."),
        ("Application API", "FastAPI and Pydantic", "Serve pages, validate inputs, expose JSON prediction, and coordinate prediction and storage."),
        ("Data processing", "pandas and MinMaxScaler", "Scale numeric fields and recreate the model's one-hot encoded input columns."),
        ("Prediction", "scikit-learn linear SVC", "Classify each transformed row as YES or NO using the saved decision boundary."),
        ("Persistence", "joblib", "Load the fitted model, scaler, column order, category options, and metadata as one bundle."),
        ("Database", "MongoDB Atlas and PyMongo", "Store original training rows and future user-response documents in separate collections."),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(table, [1900, 2200, 5260])
    format_table_text(table, font_size=8.8)
    add_caption(doc, "Table 1. Main technologies and ownership boundaries.")

    doc.add_heading("4.2 Request sequence", level=2)
    steps = [
        "The user completes profile fields and ten Yes/No questions.",
        "FastAPI parses the form and validates required values and age limits.",
        "Reverse-scoring rules convert the raw answers into ten binary AQ point values.",
        "The ten points are summed into the `result` feature.",
        "Age and result are transformed by the fitted MinMaxScaler.",
        "Categorical values are one-hot encoded and reindexed to the saved 103-column order.",
        "The saved SVC predicts the class; linear contributions are ranked for explanation.",
        "The result page is rendered and the response is saved to MongoDB when configured.",
    ]
    for step in steps:
        add_list_item(doc, step, ordered=True)


def add_data_and_methodology(doc, dataset_path):
    doc.add_heading("5. Dataset and Data Preparation", level=1)
    add_body(
        doc,
        "The project uses the UCI Autism Screening Adult dataset created by Fadi Thabtah [3]. The published dataset contains 704 instances and 21 features. The local CSV includes ten AQ score fields, demographic and history fields, an AQ result value, metadata fields, and the target class. The notebook drops rows containing missing values, leaving 702 records: 189 YES and 513 NO. This is a class distribution of approximately 26.9% YES and 73.1% NO."
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(dataset_path), width=Inches(6.2))
    set_alt_text(shape, "Dataset summary", "Cards showing 702 usable records, 189 YES records, and 513 NO records.")
    add_caption(doc, "Figure 2. Class composition after the notebook removes incomplete rows.")

    doc.add_heading("5.1 Production input features", level=2)
    add_body(
        doc,
        "The exported production model uses 18 raw inputs. It does not use every CSV column. Fields such as `age_desc` and `used_app_before` are excluded by the notebook's feature selection, while `Class/ASD` remains the prediction target. The spelling of several internal field names is preserved to remain compatible with the original dataset and saved model."
    )
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ("Group", "Fields", "Treatment")):
        cell.text = value
    feature_rows = [
        ("Numeric", "age, result", "Scaled to a common range with the fitted MinMaxScaler."),
        ("Categorical", "gender, ethnicity, jundice, austim, contry_of_res, relation", "Converted into binary one-hot indicator columns."),
        ("AQ-10", "A1_Score through A10_Score", "Binary 0/1 points produced from the Yes/No answers."),
        ("Target", "Class/ASD", "Converted from YES/NO to 1/0 for supervised learning."),
    ]
    for values in feature_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(table, [1500, 3900, 3960])
    format_table_text(table, font_size=8.9)
    add_caption(doc, "Table 2. Inputs used by the exported full-feature model.")

    doc.add_heading("5.2 AQ-10 answer scoring", level=2)
    add_body(
        doc,
        "The frontend asks the real question text rather than displaying A1 through A10. Each answer contributes either one point or zero points. Questions 2, 3, 4, 5, 6, and 9 are reverse-scored in the current implementation, so a No answer adds one point. The remaining questions add one point for Yes. The application calculates the total automatically and never asks the user to enter it."
    )
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ("Question", "Short description", "Answer worth 1 point")):
        cell.text = value
    questions = [
        ("1", "Notices small sounds", "Yes"),
        ("2", "Focuses on the whole picture", "No"),
        ("3", "Finds multitasking easy", "No"),
        ("4", "Returns quickly after interruption", "No"),
        ("5", "Reads between the lines", "No"),
        ("6", "Notices when a listener is bored", "No"),
        ("7", "Difficulty understanding character intentions", "Yes"),
        ("8", "Collects information about categories", "Yes"),
        ("9", "Reads thoughts or feelings from faces", "No"),
        ("10", "Difficulty understanding people's intentions", "Yes"),
    ]
    for values in questions:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(table, [1200, 5560, 2600])
    format_table_text(table, font_size=8.7)
    add_caption(doc, "Table 3. Binary AQ-10 scoring rules implemented by the browser form.")

    doc.add_heading("5.3 Preprocessing", level=2)
    for item in [
        "Missing-data handling: rows missing required training inputs or the target are removed.",
        "Numeric scaling: age and the calculated total score are transformed with MinMaxScaler.",
        "Categorical encoding: pandas one-hot encoding converts text categories into binary columns.",
        "Column alignment: prediction data is reindexed to the exact 103-column order saved during training.",
        "Label encoding: target YES becomes 1 and target NO becomes 0.",
        "Split: 80% of cleaned records are used for training and 20% for testing with random_state=1.",
    ]:
        add_list_item(doc, item)

    add_callout(
        doc,
        "Data caution",
        "The class is imbalanced and the `result` total is calculated from the same AQ answers already included as individual features. This can make separation unusually easy and can produce performance that looks stronger than real-world clinical generalization.",
        color=GOLD,
        fill="FFF7E8",
    )


def add_methodology(doc):
    doc.add_heading("6. Machine-Learning Methodology", level=1)
    add_body(
        doc,
        "The notebook compares six supervised classifiers to cover different learning assumptions. This comparative step is important because a single high score does not show whether the chosen method is uniquely suitable or whether the dataset is easy for many algorithms."
    )
    methods = [
        ("Decision Tree", "Builds understandable if/then splits. It can fit nonlinear patterns but may overfit small datasets."),
        ("Random Forest", "Combines multiple trees to reduce the instability of a single tree. The notebook uses five estimators."),
        ("Support Vector Machine", "Finds a separating boundary with maximum margin. A linear kernel also supports contribution-based explanation."),
        ("K-Nearest Neighbors", "Predicts from nearby examples. It is simple but sensitive to scaling, distance, and the choice of k."),
        ("Multinomial Naive Bayes", "Uses conditional-independence assumptions. It is fast but those assumptions do not match all mixed profile features."),
        ("Logistic Regression", "Provides a strong linear baseline and estimates a decision from a weighted combination of encoded inputs."),
    ]
    for name, description in methods:
        add_body(doc, f"{name}: {description}", bold_lead=f"{name}: ")

    doc.add_heading("6.1 Model selection and grid search", level=2)
    add_body(
        doc,
        "The final export uses `GridSearchCV`, which exhaustively evaluates combinations from a defined parameter grid through cross-validation [4]. The search evaluates C values 1 through 5, linear/poly/RBF/sigmoid kernels, and polynomial degrees 1 through 5. It selects the setting with the highest F0.5 score. The saved result is a linear SVC with C=1 and degree=1."
    )
    add_code_block(
        doc,
        "parameters = {\n"
        "    'C': range(1, 6),\n"
        "    'kernel': ['linear', 'poly', 'rbf', 'sigmoid'],\n"
        "    'degree': range(1, 6),\n"
        "}\n"
        "grid = GridSearchCV(SVC(random_state=1), parameters, scoring=f_beta_scorer)"
    )

    doc.add_heading("6.2 Evaluation metrics", level=2)
    add_body(
        doc,
        "Accuracy measures the fraction of all predictions that are correct. Precision measures how many predicted YES cases are actually YES, while recall measures how many actual YES cases are found. F-beta combines precision and recall. The notebook uses beta=0.5, so precision receives more emphasis than recall [5]. This choice discourages false-positive YES predictions, but in a screening context it can also increase the risk of missing a person who might benefit from further evaluation."
    )
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ("Metric", "Meaning", "Project use")):
        cell.text = value
    metric_rows = [
        ("Accuracy", "Correct predictions divided by all predictions", "Overall holdout and cross-validation performance"),
        ("Precision", "Correct YES predictions divided by predicted YES cases", "Used indirectly through F0.5"),
        ("Recall", "Correct YES predictions divided by actual YES cases", "Important for missed-screen analysis"),
        ("F0.5", "Weighted harmonic mean favoring precision", "Grid-search selection score"),
    ]
    for values in metric_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(table, [1700, 4000, 3660])
    format_table_text(table, font_size=8.9)
    add_caption(doc, "Table 4. Evaluation metrics and their role in this project.")

    doc.add_heading("6.3 Model persistence", level=2)
    add_body(
        doc,
        "After training, `joblib.dump` serializes a dictionary containing the fitted SVC, fitted scaler, raw feature list, numeric and categorical feature lists, encoded-column order, category choices, and a model name. FastAPI deserializes this bundle once when the process starts. This is model persistence: training happens deliberately in `export_model.py`, while production requests perform only transformation and inference."
    )


def add_implementation(doc):
    doc.add_heading("7. Implementation Details", level=1)
    doc.add_heading("7.1 Backend and routes", level=2)
    add_body(
        doc,
        "`app.py` creates the FastAPI application and loads `models/autism_full_model.joblib`. The GET `/` route renders the questionnaire. POST `/predict` accepts browser form data, performs reverse scoring, predicts, attempts database storage, and renders the result modal. POST `/api/predict` accepts a validated JSON body for programmatic use. GET `/assets/image.jpeg` serves the project image. FastAPI was selected because it integrates typed Pydantic validation and automatically exposes OpenAPI-based Swagger UI and ReDoc documentation [6]."
    )
    doc.add_heading("7.2 Input validation", level=2)
    add_body(
        doc,
        "The Pydantic `ScreeningAnswers` model restricts age to 1 through 120 and restricts each AQ score to 0 or 1. The browser route also checks that every required field exists and that raw question values are Yes or No. Invalid browser submissions return the same form with an error and preserve the values that can be reused."
    )
    doc.add_heading("7.3 Frontend experience", level=2)
    for item in [
        "Real AQ-10 statement text is displayed instead of internal column names.",
        "Yes/No radio controls provide consistent width, selected state, keyboard behavior, and hover feedback.",
        "Profile inputs use dataset-derived category options so submitted values match trained categories.",
        "A responsive modal opens after prediction and leaves the completed form behind it for editing.",
        "The modal uses simple phrases such as 'What influenced this result' and '1 point' rather than exposing model jargon.",
        "The project image is used subtly in the frontend and provides a consistent visual identity.",
    ]:
        add_list_item(doc, item)

    doc.add_heading("7.4 Result explanation", level=2)
    add_body(
        doc,
        "The selected SVC has a linear coefficient for each encoded input. For a submitted row, the application multiplies each transformed feature value by its coefficient. Positive contributions point toward YES and negative contributions point toward NO. The strongest factors in the predicted direction and the strongest opposing factors are shown separately. This is a local model calculation, not a causal explanation. The modal also lists every answer, the scoring rule, and whether that answer produced zero or one point."
    )
    add_callout(
        doc,
        "Interpretation rule",
        "A factor described as 'points toward YES' means only that it moved this linear model's decision in that direction. It does not mean the factor causes autism or that it should be used alone for medical decisions.",
        color=GOLD,
        fill="FFF7E8",
    )

    doc.add_heading("7.5 MongoDB integration", level=2)
    add_body(
        doc,
        "`database.py` reads connection settings from `.env`, creates a reusable PyMongo client, and writes new records to `user_responses`. The database name defaults to `autism_detection`; original rows can be held in `training_data`. Each response stores profile values, scored answers, raw browser answers when available, model output, explanation, model name, timestamp, source, and limited request metadata. PyMongo is MongoDB's official Python driver [7]."
    )
    add_body(
        doc,
        "Database failure is deliberately non-fatal. `save_user_response` catches PyMongo errors and returns False, allowing the result to be shown with a message that the answers could not be saved. This supports a more reliable live demonstration and avoids making screening dependent on network availability."
    )
    doc.add_heading("7.6 Future-training boundary", level=2)
    add_body(
        doc,
        "New MongoDB responses are not automatically used for training. A response contains the model's prediction, not a verified clinical outcome. Treating that prediction as a new target would cause self-training feedback and reinforce existing mistakes. Future data must be reviewed, consented, cleaned, de-identified where appropriate, and paired with a reliable label before supervised retraining."

    )


def add_benchmarking(doc, benchmark_path):
    doc.add_heading("8. Benchmarking and Evaluation", level=1)
    add_body(
        doc,
        "Benchmarking is mandatory under the project guidelines. The notebook reports 10-fold cross-validation means for five methods and holdout results for all six. The Decision Tree section reports a perfect 141-case holdout confusion matrix but does not report a cross-validation mean. The optimized SVC export reports 1.0000 test accuracy. Values below are transcribed from the notebook outputs; no missing score has been invented."
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(benchmark_path), width=Inches(6.35))
    set_alt_text(shape, "Benchmark chart", "Horizontal bars comparing cross-validation means and F0.5 holdout scores across six classifiers.")
    add_caption(doc, "Figure 3. Notebook benchmark results across six classifier families.")

    table = doc.add_table(rows=1, cols=5)
    headers = ("Method", "Configuration", "10-fold CV mean", "Holdout F0.5", "Additional holdout result")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    results = [
        ("Decision Tree", "Default", "Not reported", "1.000*", "Accuracy 1.000; confusion matrix TN=101, TP=40"),
        ("Random Forest", "5 trees, random_state=1", "0.9901", "1.0000", "-"),
        ("SVM", "Linear, C=1, gamma=2", "1.0000", "1.0000", "-"),
        ("KNN", "k=10", "0.9459", "0.9184", "Best explored k=19 CV mean 0.9573"),
        ("Naive Bayes", "Multinomial", "0.8746", "0.7675", "-"),
        ("Logistic Regression", "Default", "0.9971", "0.9949", "-"),
        ("Grid-searched SVC", "Linear, C=1, degree=1", "Selected by 5-fold grid search", "1.0000", "Test accuracy 1.0000"),
    ]
    for values in results:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(table, [1550, 2050, 1600, 1400, 2760])
    format_table_text(table, font_size=7.9)
    add_caption(doc, "Table 5. Comparative results recorded in the notebook. *F0.5 follows from the displayed perfect confusion matrix.")

    doc.add_heading("8.1 Interpretation", level=2)
    add_body(
        doc,
        "The linear SVM, Random Forest, Decision Tree, and Logistic Regression all perform extremely well. KNN is lower, and Multinomial Naive Bayes performs worst among the compared methods. The broad success of several classifiers suggests that the processed dataset is strongly separable. The linear SVC is a reasonable production choice because it matches the best reported score, works well with scaled and one-hot encoded data, and exposes a coefficient vector that supports the implemented explanation."
    )
    add_body(
        doc,
        "The grid search improves the notebook's unoptimized SVC from 0.9645 accuracy and 0.9574 F0.5 to 1.0000 for both measures. However, the same test split is repeatedly examined during notebook development, and the final export fits the scaler before the split. Both practices can make performance estimates optimistic. A stronger experiment would use a preprocessing Pipeline fitted only inside each training fold and retain a final untouched test set."
    )
    doc.add_heading("8.2 Error analysis", level=2)
    add_body(
        doc,
        "The displayed Decision Tree confusion matrix contains 101 true negatives, 40 true positives, zero false positives, and zero false negatives. Because the optimized model also reports a perfect test score, there are no test errors available for case-level analysis in the notebook. This absence of observed errors should not be read as evidence that real-world errors cannot occur. A useful next evaluation should test external or temporally separated data, inspect false-positive and false-negative subgroups, and report uncertainty by demographic category."
    )


def add_results_limitations_future(doc):
    doc.add_heading("9. Results and Discussion", level=1)
    add_body(
        doc,
        "The project produces a working end-to-end AI system rather than leaving the model inside a notebook. A user can complete the questionnaire, receive a result without waiting for training, understand how answers were scored, view important model factors, and continue even when database storage is unavailable. The same prediction logic is available through a JSON API, and FastAPI provides interactive API documentation."
    )
    achievements = [
        "Model reuse: the joblib bundle removes repeated training from startup and request handling.",
        "Feature consistency: production transformation follows the saved scaler and column order.",
        "Usability: internal A1-A10 names are replaced with readable questionnaire statements.",
        "Explainability: linear contributions and AQ scoring are presented in a structured modal.",
        "Resilience: MongoDB saving can fail without preventing prediction.",
        "Data collection: responses are stored in a structure suitable for later review and labeling.",
        "Responsible wording: the frontend and report distinguish screening from diagnosis.",
    ]
    for item in achievements:
        add_list_item(doc, item)

    doc.add_heading("9.1 Functional evaluation cases", level=2)
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ("Case", "Input/condition", "Expected behavior", "Observed project behavior")):
        cell.text = value
    cases = [
        ("Complete browser form", "All profile fields and ten answers", "Render prediction and explanation", "Result modal renders with score and factors"),
        ("Invalid or missing field", "Incomplete form or invalid age", "Reject request clearly", "Form returns HTTP 400 with completion message"),
        ("MongoDB unavailable", "No URI, blocked network, or driver error", "Keep prediction usable", "Prediction is shown; save flag is false"),
        ("JSON API", "Valid scored 0/1 AQ inputs", "Return structured result", "Prediction, score, explanation, and save status returned"),
        ("Model missing", "Joblib path absent", "Fail early with actionable message", "Startup raises instruction to export model once"),
    ]
    for values in cases:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(table, [1500, 2300, 2800, 2760])
    format_table_text(table, font_size=8.1)
    add_caption(doc, "Table 6. Functional cases used to evaluate the integrated application.")

    doc.add_heading("10. Limitations and Responsible AI", level=1)
    limitations = [
        "Small dataset: 702 cleaned records are not enough to establish broad clinical reliability.",
        "Class imbalance: approximately 73% of cleaned records are NO, so accuracy alone can be misleading.",
        "Potential leakage/shortcut: the AQ total score is derived from the ten AQ items and is supplied alongside them.",
        "Preprocessing before split: the scaler is fitted before the train/test split in the export script, which leaks range information from test data.",
        "Limited external validity: results come from one dataset and one random split, with no independent clinical test set.",
        "Simplified questionnaire: the interface reduces a multi-level questionnaire response into Yes/No choices to match the dataset representation.",
        "Sensitive features: demographic, jaundice, and family-history values can reflect correlations and bias rather than causal evidence.",
        "Explanation limits: coefficient contributions show model arithmetic, not medical reasons or causal mechanisms.",
        "Unverified future labels: stored predictions cannot be treated as correct outcomes for retraining.",
        "Privacy: health-related answers and request metadata require consent, access control, retention limits, and secure credential management.",
    ]
    for item in limitations:
        add_list_item(doc, item)
    add_callout(
        doc,
        "Safety boundary",
        "The system should be described as an educational screening prototype. It must not be marketed as diagnosing autism, and a YES or NO result must not be used to deny care, services, education, employment, or other opportunities.",
        color=RED,
        fill="FCEEEF",
    )

    doc.add_heading("11. Future Improvements", level=1)
    future_items = [
        "Use a scikit-learn Pipeline so imputation, scaling, and one-hot encoding are fitted only on training folds.",
        "Reserve a final untouched test set or evaluate on an external dataset from a different population.",
        "Report recall, specificity, ROC-AUC, PR-AUC, calibration, and subgroup performance in addition to accuracy and F0.5.",
        "Reconsider metric priorities for screening; compare F1 or F2 when missed positives are more costly than false alarms.",
        "Test a model without the derived `result` feature to measure how much it acts as a shortcut.",
        "Add model version, data version, preprocessing version, and consent status to every stored response.",
        "Remove or hash request metadata unless it is necessary and explicitly disclosed.",
        "Create an expert-reviewed labeling workflow before including future responses in training.",
        "Add automated unit tests, API integration tests, browser accessibility tests, and deployment health checks.",
        "Conduct usability testing with autistic adults and clinicians, then revise wording and interaction design based on feedback.",
        "Add secure authentication and role-based database access for any administrative or retraining interface.",
        "Monitor data drift and performance after deployment rather than assuming notebook scores remain valid."
    ]
    for item in future_items:
        add_list_item(doc, item)

    doc.add_heading("12. Conclusion", level=1)
    add_body(
        doc,
        "This project demonstrates the complete path from exploratory machine learning to a usable AI application. Six classifiers were compared, an SVC was tuned through grid search, and the selected full-feature model was serialized with the preprocessing contract required for consistent inference. FastAPI provides validation, browser and JSON routes, Jinja2 rendering, and automatic documentation. The frontend presents ten understandable questions and a responsive explanation modal, while MongoDB optionally stores responses for careful future review."
    )
    add_body(
        doc,
        "The strongest technical lesson is that deployment is more than calling `predict()`. The application must reproduce training transformations, validate user input, preserve model versions, handle unavailable services, explain outputs, and communicate limitations. The reported model scores are high, but responsible interpretation requires attention to data leakage, class imbalance, sample size, external validity, privacy, and the fact that screening is not diagnosis. The system is therefore best understood as an educational AI implementation and a foundation for more rigorous evaluation, not as a finished clinical tool."
    )


def add_references_and_appendix(doc):
    doc.add_page_break()
    doc.add_heading("References", level=1)
    references = [
        ("[1] Allison, C., Auyeung, B., and Baron-Cohen, S. (2012). Toward brief red flags for autism screening: The short Autism Spectrum Quotient and the short Quantitative Checklist in 1,000 cases and 3,000 controls. Journal of the American Academy of Child & Adolescent Psychiatry, 51(2), 202-212. DOI: 10.1016/j.jaac.2011.11.003.", "https://doi.org/10.1016/j.jaac.2011.11.003"),
        ("[2] Ashwood, K. L. et al. (2016). Predicting the diagnosis of autism in adults using the Autism-Spectrum Quotient questionnaire. Psychological Medicine, 46(12), 2595-2604.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4988267/"),
        ("[3] Thabtah, F. (2017). Autism Screening Adult. UCI Machine Learning Repository. DOI: 10.24432/C5F019.", "https://archive.ics.uci.edu/dataset/426/autism+screening+adult"),
        ("[4] scikit-learn developers. GridSearchCV documentation.", "https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.GridSearchCV.html"),
        ("[5] scikit-learn developers. F-beta score documentation.", "https://scikit-learn.org/stable/modules/generated/sklearn.metrics.fbeta_score.html"),
        ("[6] FastAPI. Framework features and automatic OpenAPI documentation.", "https://fastapi.tiangolo.com/features/"),
        ("[7] MongoDB. PyMongo Driver documentation and FastAPI integration guidance.", "https://www.mongodb.com/docs/languages/python/pymongo-driver/current/"),
        ("[8] Project source files: autism_detection.ipynb, export_model.py, app.py, database.py, templates/index.html, and static/styles.css. Local implementation, June 2026.", None),
    ]
    for text, url in references:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(8)
        run = paragraph.add_run(text)
        set_run_font(run, size=9.5)
        if url:
            paragraph.add_run(" ")
            add_hyperlink(paragraph, url, url)

    doc.add_heading("Appendix A. Project File Guide", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "File"
    table.rows[0].cells[1].text = "Purpose"
    files = [
        ("autism_detection.ipynb", "Exploration, preprocessing, six-model comparison, grid search, and export cell."),
        ("autism_data.csv", "Original adult autism screening dataset."),
        ("export_model.py", "Repeatable one-time training and serialization of the production bundle."),
        ("models/autism_full_model.joblib", "Saved SVC and preprocessing metadata loaded by FastAPI."),
        ("app.py", "Question definitions, validation, transformation, prediction, explanation, and routes."),
        ("database.py", "MongoDB client, response insertion, and original-row upsert helpers."),
        ("templates/index.html", "Questionnaire, result modal, accessible labels, and modal behavior."),
        ("static/styles.css", "Responsive visual design, radio states, result modal, and image treatment."),
        (".env / .env.example", "Local database secrets and a safe configuration template."),
        ("README.md", "Setup, architecture, model details, API usage, and troubleshooting."),
    ]
    for values in files:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = values
    set_table_geometry(table, [3050, 6310])
    format_table_text(table, font_size=8.8)
    add_caption(doc, "Table A1. Project files and their responsibilities.")

    doc.add_heading("Appendix B. Run and Demonstration Guide", level=1)
    add_code_block(doc, "python -m uvicorn app:app --reload --host 127.0.0.1 --port 8000")
    add_body(doc, "Open http://127.0.0.1:8000 in a browser. If port 8000 is blocked, replace it with 8010.")
    demo_steps = [
        "Show the home page and identify profile fields and the ten real questionnaire statements.",
        "Complete every field and explain that reverse scoring converts raw Yes/No choices into 0/1 points.",
        "Submit the form and show the result modal, AQ total, profile summary, influential details, and per-answer scoring.",
        "Close the modal to demonstrate that the completed form remains available for editing.",
        "Show `/docs` and the POST `/api/predict` endpoint.",
        "Show a successful MongoDB save, or explain the save-status message if the database is unavailable.",
        "Demonstrate one limitation: disconnect MongoDB or use an invalid input and show that the system responds safely.",
        "Explain that the joblib file is loaded once and that new responses do not automatically retrain the model.",
    ]
    for step in demo_steps:
        add_list_item(doc, step, ordered=True)

    doc.add_heading("Appendix C. Example API Request", level=1)
    add_code_block(
        doc,
        "POST /api/predict\n"
        "Content-Type: application/json\n\n"
        "{\n"
        "  \"age\": 24,\n"
        "  \"gender\": \"m\",\n"
        "  \"ethnicity\": \"White-European\",\n"
        "  \"jundice\": \"no\",\n"
        "  \"austim\": \"no\",\n"
        "  \"contry_of_res\": \"India\",\n"
        "  \"relation\": \"Self\",\n"
        "  \"A1_Score\": 1, \"A2_Score\": 0, \"A3_Score\": 1,\n"
        "  \"A4_Score\": 0, \"A5_Score\": 1, \"A6_Score\": 0,\n"
        "  \"A7_Score\": 1, \"A8_Score\": 0, \"A9_Score\": 1,\n"
        "  \"A10_Score\": 0\n"
        "}"
    )
    add_body(
        doc,
        "For direct API use, A1_Score through A10_Score should be sent as already-scored 0/1 values. The browser form performs reverse scoring automatically. The response contains the prediction, calculated score, model inputs, explanation, and `saved_to_mongodb` status."
    )

    doc.add_heading("Appendix D. Q&A Preparation", level=1)
    qa = [
        ("Why was this topic selected?", "It combines a meaningful screening problem with the complete AI lifecycle: data preparation, model comparison, tuning, persistence, API integration, explanation, storage, and responsible-use analysis."),
        ("Why use SVC?", "The tuned linear SVC matched the best reported benchmark, handles scaled one-hot features well, and exposes linear coefficients for the implemented explanation."),
        ("Why not retrain at startup?", "Training is slower, less predictable, and unnecessary for every request. Model persistence gives consistent inference and makes deployment simpler."),
        ("Why use F0.5?", "The notebook's metric gives more weight to precision than recall. This reduces false-positive emphasis, although the report notes that screening may require reconsidering this trade-off."),
        ("Why can 100% accuracy be misleading?", "The dataset is small and strongly separable, preprocessing occurs before splitting, the derived total score overlaps with question inputs, and there is no independent external test set."),
        ("Are MongoDB responses ready for training?", "No. They contain model predictions, not verified outcomes. They require consent, review, cleaning, reliable labels, and a new evaluation before use."),
        ("What does the explanation prove?", "It shows which transformed values moved the linear decision toward YES or NO. It does not establish causes or provide a medical explanation."),
        ("What is the most important next improvement?", "Build a leakage-resistant preprocessing Pipeline and evaluate it on independent data with recall, specificity, calibration, and subgroup analysis."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Likely question"
    table.rows[0].cells[1].text = "Answer to understand"
    for values in qa:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = values
    set_table_geometry(table, [3300, 6060])
    format_table_text(table, font_size=8.4)
    add_caption(doc, "Table A2. Questions aligned with the course evaluation guidance.")


def build_report():
    architecture_path = WORK_DIR / "architecture.png"
    benchmark_path = WORK_DIR / "benchmarks.png"
    dataset_path = WORK_DIR / "dataset_summary.png"
    create_architecture_diagram(architecture_path)
    create_benchmark_chart(benchmark_path)
    create_dataset_chart(dataset_path)

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_problem_and_background(doc)
    add_architecture(doc, architecture_path)
    add_data_and_methodology(doc, dataset_path)
    add_methodology(doc)
    add_implementation(doc)
    add_benchmarking(doc, benchmark_path)
    add_results_limitations_future(doc)
    add_references_and_appendix(doc)

    core = doc.core_properties
    core.title = "Autism Detection Using Machine Learning - Project Report"
    core.subject = "Introduction to AI and Generative AI final system implementation"
    core.author = "Student Project"
    core.keywords = "autism screening, AQ-10, machine learning, FastAPI, SVC, MongoDB"
    core.comments = "Generated from the verified project implementation and supplied course guidelines."
    doc.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    build_report()
